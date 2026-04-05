"""
Service 2: Welcome Letter Service
Owner: Nawaf

Purpose:
- Receive validated welcome-letter records from Service 1
- Open the welcome letter template
- Replace placeholders with customer information
- Save personalized welcome letters
"""

import os
import re
from pathlib import Path

from docx import Document
from flask import Flask, jsonify, request

app = Flask(__name__)

PORT = int(os.getenv('PORT', '5001'))
SERVICE_NAME = os.getenv('SERVICE_NAME', 'welcome-service')
TEMPLATE_FILE = Path(os.getenv('TEMPLATE_FILE', '/app/template.docx'))
OUTPUT_DIR = Path(os.getenv('OUTPUT_DIR', '/app/shared/output/letters'))


def safe_filename(text):
    """Make a string safe to use in a file name."""
    return re.sub(r'[^A-Za-z0-9_-]', '_', str(text))


def replace_placeholders(document, customer):
    """Replace template placeholders in paragraphs and tables."""
    # Replace in normal paragraphs
    for paragraph in document.paragraphs:
        new_text = paragraph.text
        for key, value in customer.items():
            new_text = new_text.replace('{{' + key + '}}', str(value))
        if new_text != paragraph.text:
            paragraph.text = new_text

    # Replace inside tables too, in case the template uses them
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                new_text = cell.text
                for key, value in customer.items():
                    new_text = new_text.replace('{{' + key + '}}', str(value))
                if new_text != cell.text:
                    cell.text = new_text


def generate_welcome_letter(customer):
    """Create one welcome letter from the template."""
    if not TEMPLATE_FILE.exists():
        raise FileNotFoundError(f'Template file not found: {TEMPLATE_FILE}')

    # Step 1: Open the welcome template
    document = Document(str(TEMPLATE_FILE))

    # Step 2: Replace placeholders with customer data
    replace_placeholders(document, customer)

    # Step 3: Create the output folder if needed
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    # Step 4: Save the final welcome letter
    file_name = (
        f"welcome_{safe_filename(customer.get('ACCOUNT_NUMBER', 'unknown'))}_"
        f"{safe_filename(customer.get('LAST_NAME', 'customer'))}_"
        f"{safe_filename(customer.get('FIRST_NAME', 'unknown'))}.docx"
    )
    output_path = OUTPUT_DIR / file_name
    document.save(output_path)

    return str(output_path)


@app.get('/health')
def health():
    """Simple health check route."""
    return jsonify({'status': 'ok', 'service': SERVICE_NAME})


@app.post('/generate')
def generate():
    """Receive validated welcome records and generate letters."""
    data = request.get_json(silent=True) or {}
    records = data.get('records', [])

    generated_files = []
    skipped_records = []

    for customer in records:
        # Only process valid welcome records
        if customer.get('LETTER_TYPE') != 'welcome' or customer.get('do_not_process'):
            skipped_records.append({
                'ACCOUNT_NUMBER': str(customer.get('ACCOUNT_NUMBER', 'unknown')),
                'reason': 'Record is not a valid welcome letter record.'
            })
            continue

        output_file = generate_welcome_letter(customer)
        generated_files.append(output_file)

    return jsonify({
        'service': SERVICE_NAME,
        'generated_count': len(generated_files),
        'generated_files': generated_files,
        'skipped_count': len(skipped_records),
        'skipped_records': skipped_records
    })


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=PORT, debug=False)
